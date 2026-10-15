"""Renderer first (acceptance rule 1): every layout rule against a Fernhill sample.

Cover, header, footer, numbering, palette, tables, a given Word template,
traceability, determinism, and the G5 guard.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document as Word
from docx.oxml.ns import qn

from procgov.config import load_config
from procgov.render.docx import OutputInsideLibrary, render_document


def _doc(library, doc_id):
    return next(d for d in library.documents if d.id == doc_id)


def _render(config, library, tmp_path, doc_id="PR-001"):
    return Word(
        str(render_document(_doc(library, doc_id), config, tmp_path / "out" / f"{doc_id}.docx"))
    )


def _field_codes(element) -> list[str]:
    return [t.text.strip() for t in element.iter(qn("w:instrText"))]


def _texts(element) -> str:
    return "".join(t.text or "" for t in element.iter(qn("w:t")))


# -- cover -------------------------------------------------------------------


def test_cover_has_title_subtitle_metadata_and_a_page_break(config, library, tmp_path):
    word = _render(config, library, tmp_path)
    styles = [(p.style.name, p.text) for p in word.paragraphs[:2]]
    assert styles == [
        ("Title", "Seed Ordering"),
        ("Subtitle", "How the shop orders seed stock each season"),
    ]
    cover = word.tables[0]
    rows = [[c.text for c in r.cells] for r in cover.rows]
    assert rows == [
        ["Document ID", "PR-001"],
        ["Type", "Process"],
        ["Status", "DOCUMENT"],
        ["Owner", "Head of Purchasing"],
    ]
    after = cover._tbl.getnext()
    assert after.find(".//" + qn("w:br")).get(qn("w:type")) == "page"


# -- header and footer ---------------------------------------------------------


def test_header_has_wordmark_label_and_document_title_field(config, library, tmp_path):
    word = _render(config, library, tmp_path)
    header = word.sections[0].header._element
    assert "Fernhill Process Library" in _texts(header)
    assert 'STYLEREF "Title"' in _field_codes(header)
    assert header.findall(".//" + qn("w:drawing")), "wordmark image missing"


def test_footer_has_classification_page_x_of_y_and_copyright(config, library, tmp_path):
    word = _render(config, library, tmp_path)
    footer = word.sections[0].footer._element
    text = _texts(footer)
    assert "Document classification: Internal" in text
    assert "Copyright Fernhill Garden Supplies, All rights reserved" in text
    assert _field_codes(footer) == ["PAGE", "NUMPAGES"]


# -- numbering -----------------------------------------------------------------


def test_sections_are_unnumbered_by_default(config, library, tmp_path):
    word = _render(config, library, tmp_path)
    headings = [p.text for p in word.paragraphs if p.style.name == "Heading 1"]
    assert headings == ["Purpose", "Scope", "Roles", "Steps"]


def test_number_headings_numbers_the_sections(fernhill, tmp_path):
    cfg = fernhill / "procgov.yaml"
    cfg.write_text(cfg.read_text() + "number_headings: true\n")
    config = load_config(cfg)
    from procgov.model import load_library

    word = _render(config, load_library(config.library), tmp_path)
    headings = [p.text for p in word.paragraphs if p.style.name == "Heading 1"]
    assert headings[0].startswith("1") and headings[3].startswith("4")


def test_numbered_steps_stay_a_numbered_list(config, library, tmp_path):
    word = _render(config, library, tmp_path)
    steps = [p for p in word.paragraphs if p.text.startswith("The shop manager reports")]
    assert steps and steps[0]._p.find(".//" + qn("w:numPr")) is not None


# -- palette -------------------------------------------------------------------


def test_headings_and_title_use_the_configured_palette(config, library, tmp_path):
    word = _render(config, library, tmp_path)
    assert str(word.styles["Heading 1"].font.color.rgb) == "3B7A3B"
    assert str(word.styles["Title"].font.color.rgb) == "3B7A3B"
    assert str(word.styles["Heading 2"].font.color.rgb) == "1E3D1E"


def test_table_header_row_is_shaded_with_the_light_colour(config, library, tmp_path):
    word = _render(config, library, tmp_path)
    style = word.styles["Table"].element
    fills = [s.get(qn("w:fill")) for s in style.iter(qn("w:shd"))]
    assert "D5E8D4" in fills


def test_bullets_use_the_primary_colour(config, library, tmp_path):
    word = _render(config, library, tmp_path)
    numbering = word.part.numbering_part.element
    colours = {
        lvl.find(qn("w:rPr")).find(qn("w:color")).get(qn("w:val"))
        for lvl in numbering.iter(qn("w:lvl"))
        if lvl.find(qn("w:numFmt")) is not None
        and lvl.find(qn("w:numFmt")).get(qn("w:val")) == "bullet"
    }
    assert colours == {"3B7A3B"}


# -- tables --------------------------------------------------------------------


def test_body_table_rows_do_not_split_and_the_header_repeats(config, library, tmp_path):
    word = _render(config, library, tmp_path)
    roles = word.tables[1]
    first = roles.rows[0]._tr.trPr
    assert first.find(qn("w:tblHeader")) is not None
    assert all(r._tr.trPr.find(qn("w:cantSplit")) is not None for r in roles.rows)


def test_wide_tables_get_a_smaller_font(config, library, tmp_path):
    doc = _doc(library, "PR-001")
    wide = "| " + " | ".join(f"C{i}" for i in range(7)) + " |\n|" + "---|" * 7 + "\n"
    wide += "| " + " | ".join("x" for _ in range(7)) + " |\n"
    doc.path.write_text(doc.path.read_text() + "\n" + wide)
    from procgov.model import load_library

    word = _render(config, load_library(config.library), tmp_path)
    sizes = {s.get(qn("w:val")) for s in word.tables[-1]._tbl.iter(qn("w:sz"))}
    assert sizes == {"16"}


# -- a given Word template -----------------------------------------------------


def test_layout_is_derived_from_a_given_word_template(fernhill, tmp_path):
    template = Word()
    template.styles["Heading 1"].font.color.rgb = __import__("docx").shared.RGBColor(
        0x12, 0x34, 0x56
    )
    template.sections[0].header.paragraphs[0].text = "Fernhill letterhead"
    template.save(str(fernhill / "fernhill-template.docx"))
    cfg = fernhill / "procgov.yaml"
    cfg.write_text(cfg.read_text() + "template: fernhill-template.docx\n")
    config = load_config(cfg)
    from procgov.model import load_library

    word = _render(config, load_library(config.library), tmp_path)
    assert str(word.styles["Heading 1"].font.color.rgb) == "123456"
    assert "Fernhill letterhead" in _texts(word.sections[0].header._element)


# -- traceability, determinism, G5 ---------------------------------------------


def test_the_output_names_its_source_and_hash(config, library, tmp_path):
    word = _render(config, library, tmp_path)
    doc = _doc(library, "PR-001")
    props = word.core_properties
    assert props.identifier == "PR-001"
    assert props.comments == f"source PR-001-seed-ordering.md sha256 {doc.sha256}"


def test_same_source_gives_the_same_bytes(config, library, tmp_path):
    doc = _doc(library, "PR-001")
    a = render_document(doc, config, tmp_path / "a" / "x.docx").read_bytes()
    b = render_document(doc, config, tmp_path / "b" / "x.docx").read_bytes()
    assert a == b


def test_output_inside_the_library_is_refused(config, library):
    with pytest.raises(OutputInsideLibrary):
        render_document(_doc(library, "PR-001"), config, config.library / "PR-001.docx")
    assert not (config.library / "PR-001.docx").exists()


def test_a_missing_template_is_an_error_not_a_silent_default(fernhill, tmp_path):
    cfg = fernhill / "procgov.yaml"
    cfg.write_text(cfg.read_text() + "template: missing.docx\n")
    config = load_config(cfg)
    from procgov.model import load_library

    with pytest.raises(FileNotFoundError):
        _render(config, load_library(config.library), tmp_path)


def test_work_instruction_renders_with_its_own_type_label(config, library, tmp_path):
    word = _render(config, library, tmp_path, "WI-001")
    assert [c.text for c in word.tables[0].rows[1].cells] == ["Type", "Work Instruction"]
    assert Path(tmp_path / "out" / "WI-001.docx").is_file()
