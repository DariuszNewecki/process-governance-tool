"""Build the reference .docx pandoc takes its layout from.

With a Word template in the configuration, the template is the reference:
its styles, header and footer are the organisation's layout, used as they are.
Without one, a reference is built from pandoc's default and the configured
palette, font and labels (the seed's make_reference_docx, made configurable).
"""

from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

from docx import Document as WordDocument
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor, Twips

from procgov.config import Config
from procgov.render.pandoc import pandoc_path


def _rgb(hexcolour: str) -> RGBColor:
    return RGBColor.from_string(hexcolour)


def _set_font(style, font: str, size=None, bold=None, color=None, italic=None) -> None:
    style.font.name = font
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), font)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color:
        style.font.color.rgb = _rgb(color)


def _spacing(style, before: float, after: float) -> None:
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)


def _borders(el_pr, tag: str, sides, color: str, val: str = "single", sz: int = 4) -> None:
    box = el_pr.find(qn(tag))
    if box is None:
        box = OxmlElement(tag)
        el_pr.append(box)
    for side in sides:
        edge = OxmlElement(f"w:{side}")
        if val == "none":
            edge.set(qn("w:val"), "nil")
        else:
            edge.set(qn("w:val"), val)
            edge.set(qn("w:sz"), str(sz))
            edge.set(qn("w:space"), "0")
            edge.set(qn("w:color"), color)
        box.append(edge)


def add_field(paragraph, code: str, font: str, size: float, color: str) -> None:
    """Insert a Word field (PAGE, STYLEREF ...) into ``paragraph``."""

    def run():
        r = paragraph.add_run()
        r.font.name = font
        r.font.size = Pt(size)
        r.font.color.rgb = _rgb(color)
        return r

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run()._r.append(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {code} "
    run()._r.append(instr)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run()._r.append(sep)
    run().text = " "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run()._r.append(end)


def _cell_text(cell, text, font, size, color, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.font.color.rgb = _rgb(color)
    r.font.bold = bold
    return p


def _plain_table(table) -> None:
    """Header/footer layout tables must not inherit the body table style."""
    tblpr = table._tbl.tblPr
    for tag in ("w:tblStyle", "w:tblLook"):
        for el in tblpr.findall(qn(tag)):
            tblpr.remove(el)
    look = OxmlElement("w:tblLook")
    for key, value in (
        ("w:val", "0000"),
        ("w:firstRow", "0"),
        ("w:lastRow", "0"),
        ("w:firstColumn", "0"),
        ("w:lastColumn", "0"),
        ("w:noHBand", "1"),
        ("w:noVBand", "1"),
    ):
        look.set(qn(key), value)
    tblpr.append(look)


def _cell_rule(table, side: str, color: str) -> None:
    others = [s for s in ("top", "left", "bottom", "right") if s != side]
    for row in table.rows:
        for cell in row.cells:
            tcpr = cell._tc.get_or_add_tcPr()
            _borders(tcpr, "w:tcBorders", others, color, val="none")
            _borders(tcpr, "w:tcBorders", [side], color)


def build_reference(config: Config, out: Path) -> Path:
    """Write the reference .docx for ``config`` to ``out`` and return it."""
    out.parent.mkdir(parents=True, exist_ok=True)
    if config.template is not None:
        if not config.template.is_file():
            raise FileNotFoundError(f"Word template not found: {config.template}")
        shutil.copyfile(config.template, out)
        return out
    subprocess.run(
        [pandoc_path(), "-o", str(out), "--print-default-data-file", "reference.docx"],
        check=True,
    )
    _style_default(config, out)
    return out


def _style_default(config: Config, out: Path) -> None:
    pal, font = config.palette, config.font
    doc = WordDocument(str(out))
    st = doc.styles
    names = {s.name for s in st}
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in names:
            _set_font(st[name], font, size=11, color=pal["ink"])
    _spacing(st["Body Text"], 0, 6)
    _set_font(st["Heading 1"], font, size=15, bold=True, color=pal["primary"])
    _spacing(st["Heading 1"], 16, 8)
    _set_font(st["Heading 2"], font, size=13, bold=True, color=pal["dark"])
    _spacing(st["Heading 2"], 12, 6)
    _set_font(st["Heading 3"], font, size=11.5, bold=True, color=pal["dark"])
    _spacing(st["Heading 3"], 10, 5)
    for h in ("Heading 4", "Heading 5", "Heading 6"):
        _set_font(st[h], font, size=11, bold=True, color=pal["dark"], italic=False)
    _set_font(st["Title"], font, size=22, bold=True, color=pal["primary"])
    _spacing(st["Title"], 0, 6)
    st["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_font(st["Subtitle"], font, size=13, bold=False, color=pal["dark"], italic=False)
    _spacing(st["Subtitle"], 0, 4)
    for el in st["Subtitle"].element.get_or_add_rPr().findall(qn("w:spacing")):
        el.getparent().remove(el)
    for name in ("Author", "Date", "Abstract"):
        if name in names:
            _set_font(st[name], font, size=10, color=pal["muted"])
    if "Hyperlink" in names:
        st["Hyperlink"].font.color.rgb = _rgb(pal["primary"])

    # Body tables: hairline grid, bold header row on the light palette colour.
    tbl = st["Table"]
    tpr = tbl.element.find(qn("w:tblPr"))
    if tpr is None:
        tpr = OxmlElement("w:tblPr")
        tbl.element.append(tpr)
    for old in tpr.findall(qn("w:tblBorders")):
        tpr.remove(old)
    _borders(
        tpr, "w:tblBorders", ["top", "left", "bottom", "right", "insideH", "insideV"], pal["line"]
    )
    mar = OxmlElement("w:tblCellMar")
    for side, width in (("top", 60), ("left", 100), ("bottom", 60), ("right", 100)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(width))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tpr.append(mar)
    for old in tbl.element.findall(qn("w:tblStylePr")):
        tbl.element.remove(old)
    first = OxmlElement("w:tblStylePr")
    first.set(qn("w:type"), "firstRow")
    frpr = OxmlElement("w:rPr")
    frpr.append(OxmlElement("w:b"))
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), pal["dark"])
    frpr.append(colour)
    first.append(frpr)
    ftc = OxmlElement("w:tcPr")
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), pal["light"])
    ftc.append(shd)
    first.append(ftc)
    tbl.element.append(first)

    # Page: A4 portrait, 2.5 cm margins.
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, side, Twips(1418))
    usable = sec.page_width - sec.left_margin - sec.right_margin

    # Header: wordmark | label + document title | (empty), rule below.
    hdr = sec.header
    hdr.is_linked_to_previous = False
    for p in list(hdr.paragraphs):
        p._p.getparent().remove(p._p)
    t = hdr.add_table(1, 3, usable)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _plain_table(t)
    _cell_rule(t, "bottom", pal["dark"])
    for i, w in enumerate((Twips(2300), Twips(4426), Twips(2300))):
        t.rows[0].cells[i].width = w
    pic = t.rows[0].cells[0].paragraphs[0]
    pic.paragraph_format.space_after = Pt(0)
    if config.wordmark is not None:
        pic.add_run().add_picture(str(config.wordmark), width=Inches(1.04))
    _cell_text(
        t.rows[0].cells[1],
        config.header_label,
        font,
        9,
        pal["dark"],
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    p2 = t.rows[0].cells[1].add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(0)
    add_field(p2, 'STYLEREF "Title"', font, 9, pal["dark"])
    hdr.add_paragraph().paragraph_format.space_after = Pt(0)

    # Footer: classification | page X of Y | copyright, rule above.
    ftr = sec.footer
    ftr.is_linked_to_previous = False
    for p in list(ftr.paragraphs):
        p._p.getparent().remove(p._p)
    f = ftr.add_table(1, 3, usable)
    _plain_table(f)
    _cell_rule(f, "top", pal["line"])
    for i, w in enumerate((Twips(3609), Twips(1808), Twips(3609))):
        f.rows[0].cells[i].width = w
    p = _cell_text(f.rows[0].cells[0], "Document classification: ", font, 8, pal["muted"])
    r = p.add_run(config.classification_label)
    r.font.name = font
    r.font.size = Pt(8)
    r.font.bold = True
    r.font.color.rgb = _rgb(pal["ok"])
    pc = f.rows[0].cells[1].paragraphs[0]
    pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pc.paragraph_format.space_after = Pt(0)
    add_field(pc, "PAGE", font, 8, pal["muted"])
    pc.add_run(" / ").font.size = Pt(8)
    add_field(pc, "NUMPAGES", font, 8, pal["muted"])
    _cell_text(
        f.rows[0].cells[2], config.copyright, font, 8, pal["muted"], align=WD_ALIGN_PARAGRAPH.RIGHT
    )
    ftr.add_paragraph().paragraph_format.space_after = Pt(0)
    doc.save(str(out))
