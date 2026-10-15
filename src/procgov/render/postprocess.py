"""What pandoc cannot do, applied to its output (the seed's docx_postprocess, extended).

- cover block: a metadata table (id, type, status, owner) under the title, then a page break;
- tables: rows never split across pages, the header row repeats; 7+ columns get 8 pt and
  tighter padding so they fit portrait A4;
- bullets in the primary palette colour;
- traceability: the document's id and source hash in the file's properties;
- determinism: fixed timestamps, so the same source gives the same bytes.
"""

from __future__ import annotations

import datetime
import io
import zipfile
from pathlib import Path

from docx import Document as WordDocument
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from procgov.config import Config
from procgov.model import Document

FIXED_TIME = datetime.datetime(2000, 1, 1, tzinfo=datetime.UTC)
_ZIP_TIME = (1980, 1, 1, 0, 0, 0)


def _flag(trpr, tag: str) -> None:
    if trpr.find(qn(tag)) is None:
        trpr.append(OxmlElement(tag))


def _finish_tables(word) -> None:
    for table in word.tables:
        if getattr(table, "_procgov_cover", False):
            continue
        for i, row in enumerate(table.rows):
            trpr = row._tr.get_or_add_trPr()
            _flag(trpr, "w:cantSplit")
            if i == 0:
                _flag(trpr, "w:tblHeader")
        if len(table.columns) >= 7:
            tblpr = table._tbl.tblPr
            for old in tblpr.findall(qn("w:tblCellMar")):
                tblpr.remove(old)
            mar = OxmlElement("w:tblCellMar")
            for side, width in (("top", 40), ("left", 50), ("bottom", 40), ("right", 50)):
                e = OxmlElement(f"w:{side}")
                e.set(qn("w:w"), str(width))
                e.set(qn("w:type"), "dxa")
                mar.append(e)
            tblpr.append(mar)
            for run in table._tbl.iter(qn("w:r")):
                rpr = run.find(qn("w:rPr"))
                if rpr is None:
                    rpr = OxmlElement("w:rPr")
                    run.insert(0, rpr)
                for tag in ("w:sz", "w:szCs"):
                    for old in rpr.findall(qn(tag)):
                        rpr.remove(old)
                    e = OxmlElement(tag)
                    e.set(qn("w:val"), "16")
                    rpr.append(e)


def _colour_bullets(word, colour: str) -> None:
    try:
        numbering = word.part.numbering_part.element
    except (KeyError, NotImplementedError):
        return
    for lvl in numbering.iter(qn("w:lvl")):
        fmt = lvl.find(qn("w:numFmt"))
        if fmt is not None and fmt.get(qn("w:val")) == "bullet":
            rpr = lvl.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                lvl.append(rpr)
            for old in rpr.findall(qn("w:color")):
                rpr.remove(old)
            c = OxmlElement("w:color")
            c.set(qn("w:val"), colour)
            rpr.append(c)


def _cover(word, doc: Document, config: Config) -> None:
    """Metadata table after the title block (Title, Subtitle), then a page break."""
    anchor = None
    for p in word.paragraphs:
        if p.style is not None and p.style.name in ("Title", "Subtitle"):
            anchor = p
        elif anchor is not None:
            break
    if anchor is None:
        return
    dtype = config.doc_type(doc.type)
    rows = (
        ("Document ID", doc.id),
        ("Type", dtype.label if dtype else doc.type),
        ("Status", doc.status),
        ("Owner", doc.owner),
    )
    table = word.add_table(rows=len(rows), cols=2)
    table._procgov_cover = True
    for (label, value), row in zip(rows, table.rows):
        for cell, text, bold in ((row.cells[0], label, True), (row.cells[1], value, False)):
            run = cell.paragraphs[0].add_run(text)
            run.font.bold = bold
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(config.palette["dark"])
    brk = word.add_paragraph()
    brk.add_run().add_break(WD_BREAK.PAGE)
    anchor._p.addnext(table._tbl)
    table._tbl.addnext(brk._p)


def _trace(word, doc: Document, config: Config) -> None:
    props = word.core_properties
    props.title = doc.title
    props.identifier = doc.id
    props.subject = config.doc_type(doc.type).label if config.doc_type(doc.type) else doc.type
    props.author = doc.owner
    props.comments = f"source {doc.rel} sha256 {doc.sha256}"
    props.created = FIXED_TIME
    props.modified = FIXED_TIME
    props.last_modified_by = "procgov"
    props.revision = 1


def normalise_zip(data: bytes) -> bytes:
    """Same content, fixed entry order and timestamps: same bytes every time."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for name in sorted(src.namelist()):
            info = zipfile.ZipInfo(name, date_time=_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            dst.writestr(info, src.read(name))
    return out.getvalue()


def postprocess(path: Path, doc: Document, config: Config) -> None:
    """Finish the pandoc output at ``path`` in place."""
    word = WordDocument(str(path))
    _cover(word, doc, config)
    _finish_tables(word)
    _colour_bullets(word, config.palette["primary"])
    _trace(word, doc, config)
    buf = io.BytesIO()
    word.save(buf)
    path.write_bytes(normalise_zip(buf.getvalue()))
